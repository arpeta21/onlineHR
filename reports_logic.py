"""
HRMS Data Export & CSV / Excel Reports Engine.
Generates downloadable report files for Employees, Attendance, Payroll, and Appraisals.
"""

import csv
from io import StringIO, BytesIO
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from models import User, EmployeeProfile, Attendance, Payslip, Appraisal, EmployeeCTC
from time_utils import today_ist


def csv_to_styled_xlsx(csv_data, sheet_name="Report"):
    """Convert report CSV text into a presentation-ready Excel workbook."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = (sheet_name or "Report")[:31]
    rows = csv.reader(StringIO(csv_data or ""))
    for row in rows:
        worksheet.append(row)

    if worksheet.max_row:
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(name="Aptos Display", size=10, bold=True, color="FFFFFF")
        body_font = Font(name="Aptos", size=10, color="1F2937")
        stripe_fill = PatternFill("solid", fgColor="F3F6FA")
        border = Border(bottom=Side(style="thin", color="D9E2F3"))
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[1].height = 32
        for row_index in range(2, worksheet.max_row + 1):
            for cell in worksheet[row_index]:
                cell.font = body_font
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=False)
                if row_index % 2 == 0:
                    cell.fill = stripe_fill
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        worksheet.print_title_rows = "1:1"
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        for column_index, column_cells in enumerate(worksheet.columns, start=1):
            values = [str(cell.value or "") for cell in column_cells]
            width = min(max(max((len(value) for value in values), default=10) + 2, 11), 34)
            worksheet.column_dimensions[get_column_letter(column_index)].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def mask_sensitive(value, keep=4):
    if value in (None, ""):
        return "-"
    text = str(value).strip()
    if not text:
        return "-"
    if len(text) <= 8:
        return "*" * len(text)
    visible = max(keep, 2)
    return f"{text[:visible]}{'*' * (len(text) - visible)}"


def generate_employee_report_csv(tenant_id=None):
    """Export all employee profiles to CSV."""
    output = StringIO()
    writer = csv.writer(output)

    writer.writerow([
        "Employee Code", "Full Name", "Role", "Designation", "Department",
        "Date of Joining", "Reporting Manager", "Founding Member", "Monthly CTC (INR)",
        "Email", "Phone", "PAN Number", "Aadhaar Number", "Bank Account", "IFSC Code"
    ])

    user_query = User.query.filter(User.role.in_(["employee", "manager", "demo_admin"]))
    if tenant_id is not None:
        user_query = user_query.filter(User.tenant_id == tenant_id)
    users = user_query.order_by(User.employee_code).all()

    for u in users:
        p = u.profile
        mgr_name = p.reporting_manager.profile.full_name if p and p.reporting_manager and p.reporting_manager.profile else "-"
        writer.writerow([
            u.employee_code,
            p.full_name if p else "-",
            u.role.capitalize(),
            p.designation if p else "-",
            p.department if p else "-",
            p.date_of_joining.strftime("%Y-%m-%d") if p and p.date_of_joining else "-",
            mgr_name,
            "Yes" if p and p.is_founding_member else "No",
            f"{p.monthly_ctc:.2f}" if p else "0.00",
            p.personal_email if p else "-",
            p.phone_number if p else "-",
            mask_sensitive(p.pan_number) if p and p.pan_number else "-",
            mask_sensitive(p.aadhaar_number) if p and p.aadhaar_number else "-",
            mask_sensitive(p.bank_account_number) if p and p.bank_account_number else "-",
            p.bank_ifsc_code if p and p.bank_ifsc_code else "-"
        ])

    output.seek(0)
    return output.getvalue()


def generate_attendance_report_csv(month=None, year=None, tenant_id=None):
    """Export monthly attendance records to CSV."""
    today = today_ist()
    month = month or today.month
    year = year or today.year

    output = StringIO()
    writer = csv.writer(output)

    writer.writerow([
        "Employee Code", "Full Name", "Date", "Clock In", "Clock Out",
        "Duration", "Status", "OT Hours", "Off Day", "Notes"
    ])

    import calendar
    last_day = calendar.monthrange(year, month)[1]
    start_d = date(year, month, 1)
    end_d = date(year, month, last_day)

    user_ids = None
    if tenant_id is not None:
        user_ids = [user.id for user in User.query.filter(User.tenant_id == tenant_id).all()] or [-1]
    record_query = Attendance.query.filter(
        Attendance.date >= start_d,
        Attendance.date <= end_d
    )
    if user_ids is not None:
        record_query = record_query.filter(Attendance.user_id.in_(user_ids))
    records = record_query.order_by(Attendance.date.asc()).all()

    for r in records:
        u = r.user
        p = u.profile if u else None
        writer.writerow([
            u.employee_code if u else "-",
            p.full_name if p else "-",
            r.date.strftime("%Y-%m-%d"),
            r.clock_in.strftime("%I:%M %p") if r.clock_in else "-",
            r.clock_out.strftime("%I:%M %p") if r.clock_out else "-",
            r.work_duration,
            r.status,
            r.ot_hours or 0.0,
            "Yes" if r.is_off_day else "No",
            r.notes or "-"
        ])

    output.seek(0)
    return output.getvalue()


def generate_payroll_report_csv(month=None, year=None, tenant_id=None):
    """Export monthly payroll & payslip details to CSV."""
    today = today_ist()
    month = month or today.month
    year = year or today.year

    output = StringIO()
    writer = csv.writer(output)

    writer.writerow([
        "Employee Code", "Full Name", "Month", "Year", "Monthly CTC Used",
        "Gross Pay", "Arrears", "Incentive / Bonus", "OT Amount",
        "Total Deductions", "Loss of Pay (LOP)", "PF Deduction",
        "ESI Deduction", "Professional Tax", "TDS Deduction", "Statutory Note",
        "Net Pay"
    ])

    payslip_query = Payslip.query.filter_by(month=month, year=year)
    if tenant_id is not None:
        payslip_query = payslip_query.filter(Payslip.tenant_id == tenant_id)
    payslips = payslip_query.all()

    for ps in payslips:
        u = ps.user
        p = u.profile if u else None
        writer.writerow([
            u.employee_code if u else "-",
            p.full_name if p else "-",
            ps.month_label(),
            ps.year,
            f"{ps.monthly_ctc_used:.2f}",
            f"{ps.gross_pay:.2f}",
            f"{ps.arrears:.2f}",
            f"{ps.incentive:.2f}",
            f"{ps.ot_amount:.2f}",
            f"{ps.total_deductions:.2f}",
            f"{ps.loss_of_pay:.2f}",
            f"{ps.pf_deduction:.2f}",
            f"{ps.esi_deduction:.2f}",
            f"{ps.professional_tax:.2f}",
            f"{ps.tds_deduction:.2f}",
            ps.statutory_note or "-",
            f"{ps.net_pay:.2f}"
        ])

    output.seek(0)
    return output.getvalue()


def generate_payroll_input_sheet_csv(month=None, year=None, tenant_id=None):
    """Export the payroll input sheet using the standard payroll column order."""
    today = today_ist()
    month = month or today.month
    year = year or today.year

    output = StringIO()
    writer = csv.writer(output)

    writer.writerow([
        "Emp ID", "Name", "Designation", "Dept", "Location", "DOJ", "Bank A/C", "IFSC", "PAN", "UAN", "ESI No",
        "Days Payable", "LOP Days", "Basic", "HRA", "Conveyance", "Special Allow", "LTA", "Bonus",
        "Reimbursements", "Gross Earnings", "PF (12% of Basic)", "ESI (0.75%)", "PT", "TDS",
        "Loan / Advance", "Total Deductions", "Net Pay", "Employer PF (12%)", "Employer ESI (3.25%)", "Gratuity (4.81%)"
    ])

    from payroll_logic import get_employee_monthly_payroll_inputs
    from payroll_logic import compute_payslip_amounts, get_active_ctc
    from statutory import esi_contribution, pf_contribution
    from models import User

    def component_amount(lines, names):
        names = {name.lower() for name in names}
        return sum(
            float(line.get("amount", 0.0) or 0.0)
            for line in lines
            if line.get("name", "").strip().lower() in names
        )

    user_query = User.query.filter(User.role.in_(["employee", "manager", "demo_admin"]))
    if tenant_id is not None:
        user_query = user_query.filter(User.tenant_id == tenant_id)

    for emp in user_query.order_by(User.employee_code).all():
        stats = get_employee_monthly_payroll_inputs(emp.id, month, year)
        p = emp.profile
        ctc = get_active_ctc(emp.id, as_of_date=date(year, month, 1))
        lines, gross, _, _ = compute_payslip_amounts(ctc.monthly_ctc if ctc else 0.0, ctc)
        basic = float(stats.get("calculated_basic", 0.0) or 0.0)
        hra = component_amount(lines, {"hra", "house rent allowance", "house rent allowance (hra)"})
        conveyance = component_amount(lines, {"conveyance", "conveyance allowance"})
        special_allowance = component_amount(lines, {"special allowance", "special allowance (residual)"})
        lta = component_amount(lines, {"lta", "leave travel allowance"})
        gross_earnings = float(stats.get("calculated_gross", gross) or gross or 0.0)
        pf = float(stats.get("default_pf", 0.0) or 0.0)
        esi = float(stats.get("default_esi", 0.0) or 0.0)
        pt = float(stats.get("default_tax", 0.0) or 0.0)
        tds = float(stats.get("default_tds", 0.0) or 0.0)
        loan_advance = float(stats.get("advance_repayment", 0.0) or 0.0)
        lop_days = float(stats.get("lop_days", 0.0) or 0.0)
        bonus = 0.0
        reimbursements = 0.0
        total_deductions = pf + esi + pt + tds + loan_advance + float(stats.get("auto_lop_amount", 0.0) or 0.0)
        net_pay = gross_earnings - total_deductions
        employer_pf = pf_contribution(basic)["employer_total"] if basic else 0.0
        employer_esi = esi_contribution(gross_earnings)["employer"] if gross_earnings else 0.0
        gratuity = basic * (15.0 / 26.0 / 12.0)
        writer.writerow([
            emp.employee_code,
            p.full_name if p else emp.employee_code,
            p.designation if p and p.designation else "-",
            p.department if p and p.department else "-",
            ", ".join(filter(None, [p.current_city if p else None, p.current_state if p else None])) or "-",
            p.date_of_joining.strftime("%Y-%m-%d") if p and p.date_of_joining else "-",
            mask_sensitive(p.bank_account_number) if p and p.bank_account_number else "-",
            p.bank_ifsc_code if p and p.bank_ifsc_code else "-",
            mask_sensitive(p.pan_number) if p and p.pan_number else "-",
            "-",
            "-",
            f"{stats.get('payable_days', 0.0):.2f}",
            f"{lop_days:.2f}",
            f"{basic:.2f}",
            f"{hra:.2f}",
            f"{conveyance:.2f}",
            f"{special_allowance:.2f}",
            f"{lta:.2f}",
            f"{bonus:.2f}",
            f"{reimbursements:.2f}",
            f"{gross_earnings:.2f}",
            f"{pf:.2f}",
            f"{esi:.2f}",
            f"{pt:.2f}",
            f"{tds:.2f}",
            f"{loan_advance:.2f}",
            f"{total_deductions:.2f}",
            f"{net_pay:.2f}",
            f"{employer_pf:.2f}",
            f"{employer_esi:.2f}",
            f"{gratuity:.2f}",
        ])

    output.seek(0)
    return output.getvalue()


def generate_appraisal_report_csv(period_type=None, year=None, tenant_id=None):
    """Export performance appraisals summary to CSV."""
    output = StringIO()
    writer = csv.writer(output)

    writer.writerow([
        "Employee Code", "Full Name", "Review Period", "Period Type", "Year",
        "Evaluator / Manager", "Status", "Self Rating (Out of 5)",
        "Manager Rating (Out of 5)", "Weighted Score %", "Self Comments", "Manager Comments"
    ])

    query = Appraisal.query
    if tenant_id is not None:
        query = query.filter(Appraisal.tenant_id == tenant_id)
    if period_type:
        query = query.filter_by(period_type=period_type)
    if year:
        query = query.filter_by(year=int(year))

    appraisals = query.order_by(Appraisal.id.desc()).all()

    for a in appraisals:
        u = a.user
        p = u.profile if u else None
        ev_p = a.evaluator.profile if a.evaluator else None
        evaluator_name = ev_p.full_name if ev_p else "-"

        writer.writerow([
            u.employee_code if u else "-",
            p.full_name if p else "-",
            a.review_period,
            a.period_type,
            a.year,
            evaluator_name,
            a.status,
            a.self_rating if a.self_rating else "-",
            a.rating if a.rating else "-",
            f"{a.weighted_score:.2f}%" if a.weighted_score is not None else "-",
            a.self_comments or "-",
            a.evaluator_comments or "-"
        ])

    output.seek(0)
    return output.getvalue()
