import os
import re
from pathlib import Path
from collections import Counter

try:
    import pandas as pd
except Exception:  # pragma: no cover - optional dependency
    pd = None

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - optional dependency
    load_workbook = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - optional dependency
    PdfReader = None

try:
    from docx import Document
except Exception:  # pragma: no cover - optional dependency
    Document = None


UPLOAD_FOLDER = os.path.join(
    os.path.abspath(os.path.dirname(__file__)),
    "instance",
    "uploads",
)
def safe_document_path(filename):
    if not filename or Path(filename).name != filename:
        raise ValueError("Invalid document path")
    upload_root = Path(UPLOAD_FOLDER).resolve()
    candidate = (upload_root / filename).resolve()
    if upload_root not in candidate.parents:
        raise ValueError("Invalid document path")
    return candidate


def _normalize_text(value):
    if not value:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def _tokenize(value):
    return [token for token in re.findall(r"[a-z0-9]+", _normalize_text(value)) if token]


def _expand_term_variants(tokens):
    variants = set()
    alias_map = {
        "ev": {"ev", "electric", "vehicle", "electric_vehicle"},
        "charger": {"charger", "chargers", "charging", "charge", "chargepoint", "charge_points", "charging_station", "charging_stations", "station", "stations"},
        "type": {"type", "types", "kind", "kinds", "category", "categories"},
        "policy": {"policy", "policies", "rule", "rules", "guideline", "guidelines"},
        "leave": {"leave", "leaves", "timeoff", "absence"},
        "casual": {"casual", "casuals", "casual_leave", "casual_leaves", "personal_leave", "personal_leaves"},
        "attendance": {"attendance", "present", "late", "timing", "timekeeping"},
    }

    for token in tokens:
        normalized = token.strip().lower()
        variants.add(normalized)
        for key, values in alias_map.items():
            if normalized in values or normalized == key:
                variants.update(values)
                break
    return variants


def _phrase_variants(tokens):
    if not tokens:
        return set()

    variants = {""}
    for token in tokens:
        expanded = _expand_term_variants([token])
        next_variants = set()
        for prefix in variants:
            for value in expanded:
                merged = " ".join(part for part in [prefix, value] if part).strip()
                if merged:
                    next_variants.add(merged)
        variants = next_variants
    return {value for value in variants if value}


def _extract_relevant_excerpt(query, text):
    if not text:
        return ""

    query_variants = _expand_term_variants(_tokenize(query))
    query_tokens = [token for token in _tokenize(query) if token not in {"as", "per", "and", "the", "of", "in", "for", "to", "on", "at", "by"}]
    sentences = [part.strip() for part in re.split(r"(?<=[.!?])\s+", text) if part.strip()]
    if not sentences:
        return text[:220]

    best_sentence = ""
    best_score = -1
    for sentence in sentences:
        sentence_tokens = _expand_term_variants(_tokenize(sentence))
        sentence_norm = _normalize_text(sentence)
        score = len(query_variants & sentence_tokens)

        if len(query_tokens) >= 2:
            for lead_length in range(min(3, len(query_tokens)), 1, -1):
                lead_tokens = query_tokens[:lead_length]
                lead_variants = _phrase_variants(lead_tokens)
                if any(variant in sentence_norm for variant in lead_variants):
                    score += lead_length * 40

        for phrase_length in range(min(4, len(query_tokens)), 1, -1):
            for start in range(0, len(query_tokens) - phrase_length + 1):
                phrase_tokens = query_tokens[start:start + phrase_length]
                phrase_variants = _phrase_variants(phrase_tokens)
                if any(variant in sentence_norm for variant in phrase_variants):
                    score += phrase_length * 10

        if score > best_score:
            best_sentence = sentence
            best_score = score

    if best_sentence:
        return best_sentence[:220]
    return text[:220]


def _read_uploaded_document_text(filename):
    if not filename:
        return ""

    try:
        file_path = safe_document_path(filename)
    except ValueError:
        return ""
    if not os.path.exists(file_path):
        return ""

    ext = os.path.splitext(filename)[1].lower()

    try:
        if ext in {".txt", ".md", ".csv"}:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
                return handle.read()

        if ext == ".pdf" and PdfReader is not None:
            reader = PdfReader(file_path)
            pages = []
            for page in reader.pages:
                text = page.extract_text() or ""
                if text:
                    pages.append(text)
            return "\n".join(pages)

        if ext == ".docx" and Document is not None:
            doc = Document(file_path)
            return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text)

        if ext in {".xlsx", ".xls"}:
            if ext == ".xlsx" and load_workbook is not None:
                workbook = load_workbook(file_path, read_only=True, data_only=True)
                sheets = []
                for sheet in workbook.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        cells = [str(value).strip() for value in row if value is not None and str(value).strip()]
                        if cells:
                            sheets.append(" ".join(cells))
                if sheets:
                    return "\n".join(sheets)
            if pd is not None:
                excel_file = pd.ExcelFile(file_path)
                rows = []
                for sheet_name in excel_file.sheet_names:
                    df = pd.read_excel(file_path, sheet_name=sheet_name)
                    for _, row in df.fillna("").astype(str).iterrows():
                        values = [str(value).strip() for value in row.tolist() if str(value).strip()]
                        if values:
                            rows.append(" ".join(values))
                if rows:
                    return "\n".join(rows)

        if ext == ".doc" and pd is not None:
            return ""

        if ext in {".png", ".jpg", ".jpeg"}:
            return ""

    except Exception:
        return ""

    return ""


def search_policy_documents(query, resources, fallback_text=""):
    cleaned_query = _normalize_text(query)
    tokens = _tokenize(query)
    query_aliases = _expand_term_variants(tokens)

    if not tokens:
        scored = []
        for resource in resources[:5]:
            scored.append({
                "resource": resource,
                "score": 0,
                "excerpt": resource.description or resource.title,
            })
        return scored

    scored = []
    for resource in resources:
        title = getattr(resource, "title", "") or ""
        description = getattr(resource, "description", "") or ""
        category = getattr(resource, "category", "") or ""
        filename = getattr(resource, "filename", "") or ""
        filename_stem = os.path.splitext(filename)[0].replace("_", " ").replace("-", " ")
        document_text = _read_uploaded_document_text(filename)
        combined = " ".join([title, description, category, fallback_text, document_text, filename_stem])
        combined_tokens = _expand_term_variants(_tokenize(combined))
        score = 0

        if cleaned_query in _normalize_text(title):
            score += 25
        if cleaned_query in _normalize_text(description):
            score += 15
        if cleaned_query in _normalize_text(filename_stem):
            score += 18
        if cleaned_query in _normalize_text(document_text):
            score += 20

        for token in tokens:
            if token in _normalize_text(title):
                score += 10
            if token in _normalize_text(description):
                score += 6
            if token in _normalize_text(category):
                score += 4
            if token in _normalize_text(filename_stem):
                score += 5
            if token in _normalize_text(document_text):
                score += 3

        alias_overlap = len(query_aliases & combined_tokens)
        if alias_overlap:
            score += alias_overlap * 8

        if not cleaned_query or score <= 0:
            continue

        excerpt = ""
        for source in [document_text, title, description, filename_stem]:
            text = str(source or "")
            if not text:
                continue
            lowered = _normalize_text(text)
            if cleaned_query in lowered:
                excerpt = text.strip()
                break
            if any(alias in _normalize_text(text) for alias in query_aliases):
                excerpt = _extract_relevant_excerpt(query, text)
                break

        if not excerpt:
            excerpt = title or description or filename_stem or category or "Relevant policy document"

        snippet = excerpt[:220]
        scored.append({
            "resource": resource,
            "score": score,
            "excerpt": snippet,
        })

    scored.sort(key=lambda item: item["score"], reverse=True)
    return scored[:5]


def _extract_exact_fact_sentence(query, text):
    if not text or not query:
        return ""

    text = str(text)
    query_tokens = [token for token in _tokenize(query) if len(token) > 2]
    if not query_tokens:
        return ""

    query_aliases = _expand_term_variants(query_tokens)
    candidates = [part.strip() for part in re.split(r"(?<=[.!?])\s+", text) if part.strip()]
    if not candidates:
        candidates = [text.strip()]

    best_sentence = ""
    best_score = -1

    for sentence in candidates:
        sentence_norm = _normalize_text(sentence)
        has_number = bool(re.search(r"\b\d+\b", sentence))
        if not has_number:
            continue

        score = 0
        if any(token in sentence_norm for token in query_tokens):
            score += 20
        if query_aliases and (query_aliases & _expand_term_variants(_tokenize(sentence))):
            score += 30

        if score > best_score:
            best_sentence = sentence
            best_score = score

    if best_score >= 40:
        return best_sentence
    return ""


def build_policy_answer(query, resources, fallback_text=""):
    if not query or not query.strip():
        return "Ask about leave, attendance, payroll, code of conduct, travel, or hiring policy and I’ll point you to the closest published resource."

    matches = search_policy_documents(query, resources, fallback_text=fallback_text)
    if not matches:
        return (
            f"I could not find a close match for '{query}' in the uploaded policy library. "
            "Please check the handbook and policies page or ask HR for the latest guideline."
        )

    best_match = matches[0]
    resource = best_match["resource"]
    excerpt = best_match["excerpt"]

    exact_fact = _extract_exact_fact_sentence(query, excerpt)
    if exact_fact:
        return f"Exact answer: {exact_fact}"

    normalized_query = _normalize_text(query)
    title_text = _normalize_text(getattr(resource, "title", ""))
    description_text = _normalize_text(getattr(resource, "description", ""))
    filename_text = _normalize_text(os.path.splitext(getattr(resource, "filename", "") or "")[0].replace("_", " ").replace("-", " "))
    resource_text = " ".join([title_text, description_text, filename_text])
    query_tokens = [token for token in _tokenize(query) if len(token) > 2]
    matching_tokens = sum(1 for token in query_tokens if token in resource_text)
    exact_coverage = (matching_tokens / len(query_tokens)) if query_tokens else 0.0

    if exact_coverage >= 0.7 or normalized_query in resource_text:
        return (
            f"The closest match is '{resource.title}' ({resource.category}). "
            f"This document appears relevant to your question: {excerpt}"
        )

    return (
        f"I found a closely related policy document, but not an exact answer for '{query}'. "
        f"Please review '{resource.title}' ({resource.category}) for the most relevant guidance: {excerpt}"
    )


def build_people_analytics(rating_rows, attendance_rows, location_rows):
    top_rated = None
    if rating_rows:
        top_rated = max(rating_rows, key=lambda row: getattr(row, "rating", 0) or 0)

    avg_rating = 0.0
    if rating_rows:
        avg_rating = sum(float(getattr(row, "rating", 0) or 0) for row in rating_rows) / len(rating_rows)

    rating_labels = {
        5: "Excellent",
        4: "Outstanding",
        3: "Good",
        2: "Average",
        1: "Poor",
    }

    def rating_band(value):
        try:
            numeric = float(value or 0)
        except (TypeError, ValueError):
            return None
        return int(numeric) if numeric in {1.0, 2.0, 3.0, 4.0, 5.0} else None

    rating_distribution = {
        label: sum(1 for row in rating_rows if rating_band(getattr(row, "rating", 0)) == score)
        for score, label in rating_labels.items()
    }
    rating_scale = []
    for score, label in rating_labels.items():
        rating_scale.append({
            "score": score,
            "label": label,
            "self_count": sum(1 for row in rating_rows if rating_band(getattr(row, "self_rating", 0)) == score),
            "manager_count": sum(1 for row in rating_rows if rating_band(getattr(row, "rating", 0)) == score),
            "potential_count": sum(1 for row in rating_rows if rating_band(getattr(row, "potential_rating", 0)) == score),
        })
    radar_max = max(
        (max(item["self_count"], item["manager_count"], item["potential_count"]) for item in rating_scale),
        default=1,
    )
    radar_vertices = [(50, 8), (90, 35), (75, 85), (25, 85), (10, 35)]

    def radar_points(attribute):
        points = []
        for item, (vertex_x, vertex_y) in zip(rating_scale, radar_vertices):
            ratio = (item[attribute] / radar_max) if radar_max else 0
            points.append((round(50 + (vertex_x - 50) * ratio, 2), round(50 + (vertex_y - 50) * ratio, 2)))
        return " ".join(f"{x}% {y}%" for x, y in points)

    radar_series = [
        {"label": "Self", "points": radar_points("self_count"), "class_name": "radar-self"},
        {"label": "Manager", "points": radar_points("manager_count"), "class_name": "radar-manager"},
        {"label": "Potential", "points": radar_points("potential_count"), "class_name": "radar-potential"},
    ]

    most_common_status = None
    if attendance_rows:
        status_counts = Counter()
        for row in attendance_rows:
            status_counts[getattr(row, "status", "Unknown")] += int(getattr(row, "count", 0) or 0)
        if status_counts:
            most_common_status = status_counts.most_common(1)[0][0]

    location_leader = None
    if location_rows:
        top_location = max(location_rows, key=lambda row: getattr(row, "count", 0) or 0)
        location_leader = getattr(top_location, "city", "Unknown")

    department_counts = Counter()
    department_total_rating = Counter()
    for row in rating_rows:
        department = getattr(row, "department", "Unassigned") or "Unassigned"
        department_counts[department] += 1
        department_total_rating[department] += float(getattr(row, "rating", 0) or 0)
    department_breakdown = [
        {
            "department": dept,
            "count": count,
            "average_rating": round(department_total_rating[dept] / count, 2) if count else 0.0,
        }
        for dept, count in sorted(department_counts.items(), key=lambda item: (-item[1], item[0]))
    ]

    manager_counts = Counter()
    manager_total_rating = Counter()
    for row in rating_rows:
        manager = getattr(row, "manager_name", "Unassigned") or "Unassigned"
        manager_counts[manager] += 1
        manager_total_rating[manager] += float(getattr(row, "rating", 0) or 0)
    manager_breakdown = [
        {
            "manager": manager,
            "count": count,
            "average_rating": round(manager_total_rating[manager] / count, 2) if count else 0.0,
        }
        for manager, count in sorted(manager_counts.items(), key=lambda item: (-item[1], item[0]))
    ]

    rating_trend = [
        {
            "label": getattr(row, "employee_name", "Employee"),
            "value": float(getattr(row, "rating", 0) or 0),
        }
        for row in sorted(rating_rows, key=lambda row: (float(getattr(row, "rating", 0) or 0), getattr(row, "employee_name", "")), reverse=True)
    ]

    descriptive_points = []
    if top_rated:
        descriptive_points.append(f"{getattr(top_rated, 'employee_name', 'The top-rated employee')} leads the team with a manager rating of {float(getattr(top_rated, 'rating', 0) or 0):.2f}.")
    if rating_rows:
        descriptive_points.append(f"The average manager rating is {avg_rating:.2f} across {len(rating_rows)} employees.")
    if most_common_status:
        descriptive_points.append(f"Attendance is most commonly recorded as {most_common_status}.")
    if location_leader:
        descriptive_points.append(f"{location_leader} is the dominant employee location cluster.")

    predictive_points = []
    if rating_rows:
        if avg_rating >= 4.0:
            predictive_points.append("At the current score level, the team has a strong probability of sustaining high performance and lower attrition risk.")
        elif avg_rating >= 3.0:
            predictive_points.append("The trend indicates moderate performance health, but a few low scorers could drag down team productivity if left unchecked.")
        else:
            predictive_points.append("The current trajectory suggests productivity and performance risk could increase without focused coaching and manager intervention.")
    if most_common_status:
        if most_common_status.lower() in {"present", "on time"}:
            predictive_points.append("Current attendance consistency is likely to support delivery stability and reduced schedule disruption.")
        else:
            predictive_points.append(f"{most_common_status} is trending as the dominant attendance pattern, which may signal a future risk of scheduling and productivity inefficiency.")
    if location_leader:
        predictive_points.append(f"{location_leader} is emerging as the main workforce concentration, which helps forecast office capacity and local support planning.")

    prescriptive_points = []
    if rating_distribution.get("Poor", 0) or rating_distribution.get("Average", 0):
        prescriptive_points.append("Prioritize coaching and manager follow-ups for employees in the low-rating group to reduce underperformance risk.")
    if most_common_status and most_common_status.lower() in {"late", "absent", "leave"}:
        prescriptive_points.append("Introduce attendance nudges, reminder workflows, and manager check-ins to address recurring delay or absence patterns.")
    if location_leader:
        prescriptive_points.append(f"Align resource planning, travel support, and local scheduling decisions around the largest workforce cluster in {location_leader}.")
    if avg_rating < 3.5:
        prescriptive_points.append("Launch a structured development plan with targeted training, milestone reviews, and recognition efforts to improve team performance.")
    else:
        prescriptive_points.append("Keep the current recognition and mentoring rhythm going to sustain high performance and engagement.")

    return {
        "top_rated_employee": getattr(top_rated, "employee_name", "No data") if top_rated else "No data",
        "average_rating": round(avg_rating, 2),
        "rating_distribution": rating_distribution,
        "rating_scale": rating_scale,
        "radar_series": radar_series,
        "attendance_top_status": most_common_status or "No data",
        "location_leader": location_leader or "No data",
        "department_breakdown": department_breakdown,
        "manager_breakdown": manager_breakdown,
        "rating_trend": rating_trend,
        "descriptive_points": descriptive_points or ["No descriptive insights are available yet."],
        "predictive_points": predictive_points or ["No predictive signals are available from the current data."],
        "prescriptive_points": prescriptive_points or ["No prescriptive action is needed until more data is available."],
    }
