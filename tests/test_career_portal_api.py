from io import BytesIO
from secrets import token_hex

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from app import app
from models import (
    Candidate,
    CandidateRecruitmentMeta,
    JobRequisition,
    Tenant,
    User,
    db,
)


API_PATH = "/api/integrations/career-portal/applications"
API_TOKEN = "test-career-portal-token"
MAX_RESUME_BYTES = 10 * 1024 * 1024


def _make_pdf():
    output = BytesIO()
    document = canvas.Canvas(output)
    document.drawString(72, 750, "Taylor Applicant")
    document.drawString(72, 730, "taylor@example.test")
    document.drawString(72, 710, "+1 555 010 2323")
    document.drawString(72, 690, "5 years experience in Python engineering")
    document.save()
    output.seek(0)
    return output


def _make_docx():
    output = BytesIO()
    document = Document()
    document.add_paragraph("Morgan DOCX Applicant")
    document.add_paragraph("morgan@example.test")
    document.add_paragraph("6 years experience in Python engineering")
    document.save(output)
    output.seek(0)
    return output


@pytest.fixture
def career_portal_environment(monkeypatch, tmp_path):
    suffix = token_hex(6).upper()
    with app.app_context():
        tenant = Tenant(company_name=f"Career Portal Tenant {suffix}", tenant_type="internal")
        foreign_tenant = Tenant(company_name=f"Other Career Tenant {suffix}", tenant_type="internal")
        db.session.add_all([tenant, foreign_tenant])
        db.session.flush()

        requester = User(
            employee_code=f"CP{suffix}",
            role="admin",
            tenant_id=tenant.id,
            must_change_password=False,
        )
        requester.set_password("CareerPortal@123")
        foreign_requester = User(
            employee_code=f"FP{suffix}",
            role="admin",
            tenant_id=foreign_tenant.id,
            must_change_password=False,
        )
        foreign_requester.set_password("ForeignPortal@123")
        db.session.add_all([requester, foreign_requester])
        db.session.flush()

        approved_requisition = JobRequisition(
            tenant_id=tenant.id,
            title="Software Engineer",
            department="Engineering",
            number_of_positions=1,
            key_skills="Python, Flask",
            requested_by_id=requester.id,
            status="Approved",
            is_manager_approved=True,
            is_admin_approved=True,
        )
        foreign_requisition = JobRequisition(
            tenant_id=foreign_tenant.id,
            title="Foreign Software Engineer",
            department="Engineering",
            number_of_positions=1,
            key_skills="Python",
            requested_by_id=foreign_requester.id,
            status="Approved",
            is_manager_approved=True,
            is_admin_approved=True,
        )
        pending_requisition = JobRequisition(
            tenant_id=tenant.id,
            title="Pending Role",
            department="Engineering",
            number_of_positions=1,
            key_skills="Python",
            requested_by_id=requester.id,
            status="Pending Admin Approval",
        )
        db.session.add_all([approved_requisition, foreign_requisition, pending_requisition])
        db.session.commit()

        monkeypatch.setenv("CAREER_PORTAL_API_TOKEN", API_TOKEN)
        monkeypatch.setenv("CAREER_PORTAL_TENANT_ID", str(tenant.id))
        monkeypatch.setitem(app.config, "UPLOAD_FOLDER", str(tmp_path))

        yield {
            "tenant": tenant,
            "foreign_tenant": foreign_tenant,
            "approved_requisition": approved_requisition,
            "foreign_requisition": foreign_requisition,
            "pending_requisition": pending_requisition,
        }

        for candidate in Candidate.query.filter(
            Candidate.requisition_id.in_([
                approved_requisition.id,
                foreign_requisition.id,
                pending_requisition.id,
            ])
        ).all():
            db.session.delete(candidate)
        db.session.delete(approved_requisition)
        db.session.delete(foreign_requisition)
        db.session.delete(pending_requisition)
        db.session.delete(requester)
        db.session.delete(foreign_requester)
        db.session.delete(tenant)
        db.session.delete(foreign_tenant)
        db.session.commit()


def _post_application(client, token, requisition_id, resume=None, **fields):
    data = {"requisition_id": str(requisition_id), **fields}
    if resume is not None:
        data["resume"] = resume
    headers = {"Authorization": f"Bearer {token}"} if token is not None else {}
    return client.post(API_PATH, data=data, headers=headers, content_type="multipart/form-data")


def test_career_portal_accepts_valid_application_and_uses_configured_tenant(career_portal_environment):
    client = app.test_client()
    requisition = career_portal_environment["approved_requisition"]
    response = _post_application(
        client,
        API_TOKEN,
        requisition.id,
        resume=(_make_pdf(), "taylor-resume.pdf"),
        name="Taylor Applicant",
        email="taylor@example.test",
        phone="+1 555 010 2323",
        experience="5",
        location="Remote",
        cover_note="Interested in the engineering role.",
        source="external_job_portal",
        tenant_id=str(career_portal_environment["foreign_tenant"].id),
    )

    assert response.status_code == 201
    assert response.is_json
    payload = response.get_json()
    candidate = db.session.get(Candidate, payload["candidate_id"])
    assert payload == {
        "success": True,
        "candidate_id": candidate.id,
        "requisition_id": requisition.id,
        "status": "created",
    }
    assert candidate.requisition_id == requisition.id
    assert candidate.tenant_id == career_portal_environment["tenant"].id
    assert candidate.full_name == "Taylor Applicant"
    assert candidate.email == "taylor@example.test"
    assert candidate.total_exp_years == 5
    assert "Applicant location: Remote" in candidate.hr_notes
    assert "Interested in the engineering role." in candidate.hr_notes
    assert candidate.resume_filename != "taylor-resume.pdf"
    assert candidate.recruitment_meta.hiring_source == "external_job_portal"
    assert API_TOKEN.encode() not in response.data


def test_career_portal_uses_existing_parser_for_docx_resume(career_portal_environment):
    response = _post_application(
        app.test_client(),
        API_TOKEN,
        career_portal_environment["approved_requisition"].id,
        resume=(_make_docx(), "morgan-resume.docx"),
    )

    assert response.status_code == 201
    candidate = db.session.get(Candidate, response.get_json()["candidate_id"])
    assert candidate.full_name == "Morgan Docx Applicant"
    assert candidate.email == "morgan@example.test"
    assert candidate.total_exp_years == 6


def test_career_portal_rejects_invalid_bearer_token(career_portal_environment):
    response = _post_application(
        app.test_client(),
        "incorrect-token",
        career_portal_environment["approved_requisition"].id,
    )

    assert response.status_code == 401
    assert response.is_json


def test_career_portal_rejects_requisition_from_another_tenant(career_portal_environment):
    response = _post_application(
        app.test_client(),
        API_TOKEN,
        career_portal_environment["foreign_requisition"].id,
    )

    assert response.status_code == 403
    assert response.is_json


def test_career_portal_rejects_unapproved_requisition(career_portal_environment):
    response = _post_application(
        app.test_client(),
        API_TOKEN,
        career_portal_environment["pending_requisition"].id,
    )

    assert response.status_code == 400
    assert response.is_json


def test_career_portal_rejects_invalid_resume_type(career_portal_environment):
    response = _post_application(
        app.test_client(),
        API_TOKEN,
        career_portal_environment["approved_requisition"].id,
        resume=(BytesIO(b"not a resume"), "resume.txt"),
    )

    assert response.status_code == 400
    assert response.is_json


def test_career_portal_rejects_empty_resume(career_portal_environment):
    response = _post_application(
        app.test_client(),
        API_TOKEN,
        career_portal_environment["approved_requisition"].id,
        resume=(BytesIO(b""), "empty-resume.pdf"),
    )

    assert response.status_code == 400
    assert response.is_json


def test_career_portal_rejects_oversized_resume(career_portal_environment):
    oversized_pdf = BytesIO(b"%PDF-1.4\n" + b"x" * (MAX_RESUME_BYTES + 1))
    response = _post_application(
        app.test_client(),
        API_TOKEN,
        career_portal_environment["approved_requisition"].id,
        resume=(oversized_pdf, "large-resume.pdf"),
    )

    assert response.status_code == 413
    assert response.is_json